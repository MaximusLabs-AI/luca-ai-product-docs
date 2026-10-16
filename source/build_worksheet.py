"""Builds the fill-in worksheet Luca's team completes: everything v1 does not document."""
import csv, re, sys
sys.path.insert(0, '/Users/krishnakaanth/luca-ai-product-docs/source')
import spec
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter

INV = '/Users/krishnakaanth/Maximus Work/clients/luca-ai/strategy/2026-10-16_luca-ai_strategy_buyer-and-ai-evaluation-parameters.csv'
OUT = sys.argv[1]
documented = {p['id']: p for p in spec.P}
EXCLUDE = {'T16','R7','R8','R10','R12','S12','CO2','CO19','T15','T2','R1','R2','CO18'}  # internal assessments or covered as contradictions
HAVE = {'✅':'Confirmed','🟡':'Partly confirmed','🛠':'In build (per your team)','🗺':'Seen in a spec or roadmap only','❌':'Not seen','❓':'Unknown to us','—':'Not assessed','':'Not assessed'}
CLAIM = {'✅':'Stated on your website','🟡':'Mentioned in passing or in legal pages','❌':'Not on your website','⚠️':'Stated inconsistently, or beyond what we have seen','—':'—','':'—'}
AREA_FIX = {'Capital-seeking brands':'Capital questions','Agencies and fractional operators':'Agency use'}

def area(sub):
    s = re.sub(r'^\d+\.\d+\s+', '', sub); s = re.sub(r'\s*[—(].*$', '', s).strip()
    return AREA_FIX.get(s, s)

rows = []
for r in csv.DictReader(open(INV)):
    if r['section'][:1] not in '345' or r['id'] in EXCLUDE: continue
    if r['id'] in documented and documented[r['id']]['status'] == 'available': continue
    w, h, c = r['weight_or_type'], r['have_or_status'], r['claimed']
    if r['id'] in documented: ask = 'Documented in v1 with a limited scope. Please confirm the scope, or widen it.'
    elif c in ('✅','⚠️') and h not in ('✅',): ask = 'Your website states this. Please confirm it is live today, and for whom.'
    else: ask = 'Do you have this? If yes: on which plans, with what limits, and since when?'
    pr = 'P1' if (w == 'H' and c in ('✅','⚠️')) or r['id'] in documented else ('P2' if w == 'H' else 'P3')
    rows.append([r['id'], area(r['subsection']), r['parameter'], {'H':'High','M':'Medium','L':'Low'}.get(w, w), HAVE.get(h, h), CLAIM.get(c, c), pr, ask])
ORDER=['U','CO','T','S','R','FD','CF','GR','OP','SU','MC','SC','EU','AG','CS','MB']
rows.sort(key=lambda x: (x[6], ORDER.index(re.match(r'[A-Z]+', x[0]).group()), int(re.sub(r'\D', '', x[0]) or 0)))

CONTRA = [
 ('Autonomous actions', 'The Agent and How it works pages say Luca executes changes (pauses campaigns, moves budgets, edits Shopify prices and discounts, sends Klaviyo flows) with autonomy levels L0–L4. The privacy page says Google access is read-only and changes nothing. On our calls Luca has been described as read-only with a human in the loop.', '/the-agent, /how-it-works, /privacy'),
 ('Capital availability', 'The home and How it works pages label capital "Coming soon"; the Capital page describes it as live, in the present tense.', '/, /how-it-works, /capital'),
 ('Capital cost, size and fee method', 'All-in cost appears as 8–12%, 10–12%, ~12%, ~10% annualised and 5–8%. Advance size appears four ways. Pricing is "dynamic per advance" on one page and "a single fixed fee" on another.', '/capital, /how-it-works, /pricing, blog'),
 ('Lender wording', 'The pricing FAQ calls Luca "a senior operator, advisor, analyst, and lender in one"; About mentions "deep pockets" from credit funds. Your positioning is "not a financing company".', '/pricing, /about'),
 ('Plan prices', 'The pricing page body, the page meta description, the Terms (per-seat price) and an older pricing script in the page source show different figures.', '/pricing, /terms-and-conditions'),
 ('Renewal and billing terms', 'The pricing FAQ says plans end automatically after 12 months; the Terms say they auto-renew monthly with notice, and that annual upfront payment is not offered while the pricing page offers annual billing.', '/pricing, /terms-and-conditions'),
 ('Trial length', 'Trials have been described as 14 days and as 30 days; the pricing page says 30.', '/pricing, calls'),
 ('Currency display', 'A script swaps the € symbol for the visitor\'s local symbol without converting the number, including on capital amounts.', '/capital, /how-it-works, /the-agent, /pricing, /contact, /careers'),
 ('Legal entity', 'Terms name Luca Labs, Inc. (Delaware); Privacy names Luca Labs (UK) Inc. (London); blog schema names Valerian Capital (UK) Ltd.', '/terms-and-conditions, /privacy, blog schema'),
 ('How Luca describes itself', '"AI operating system for e-commerce", "AI Co-Founder", "senior operator, advisor, analyst, and lender", and an old schema block calling Luca "an AI-powered personal assistant".', '/, /about, /pricing, blog schema'),
 ('Setup time and time to first insight', 'Stated as under 5 minutes, two clicks per source, within the hour, within 24 hours, the same day, and five minutes.', '/, /how-it-works, /pricing, /contact, /careers'),
 ('Data freshness', 'Stated as real time; as nightly sync with webhook deltas; and per source (15 minutes / nightly / live pull).', '/how-it-works'),
 ('Encryption and isolation', 'TLS 1.3 (How it works) vs TLS 1.2+ (Privacy); "single-tenant infrastructure" vs "separate per-customer BigQuery datasets".', '/how-it-works, /privacy'),
 ('EU data residency', '"EU data residency available" vs transfers outside the EEA under standard contractual clauses.', '/how-it-works, /privacy'),
 ('Use of customer data', '"Used only to answer your questions" vs the Terms right to use anonymised data for benchmarking.', '/pricing, Terms §9.6'),
 ('Ratings in page code', 'Blog pages carry a 4.8 / 872 rating in structured data with no visible source.', 'blog template'),
 ('Demo details', 'Demo on demo data vs "live, on your data"; 30 vs 45 minutes; reply in one vs a few business days.', '/contact, /the-agent, /capital'),
 ('Number of data sources', '"20+ sources" on How it works; we have seen about nine connected in product.', '/how-it-works'),
 ('Payback claim', 'Customers recoup the fee in the "first week" (pricing) vs the "first month" (Mastercard page).', '/pricing, /pricing-mastercard-business-bonus'),
]
QUESTIONS = [
 'Which actions can Luca execute in a connected system today (if any), on which platforms, and with which write permissions?',
 'Do proactive alerts and the daily insight run automatically today? Which alert types, and through which channels (app, email, Slack, WhatsApp)?',
 'Does Luca keep memory across conversations? Can a user see, edit or delete it?',
 'Which connectors are live in production today? Please list every one (e.g. Klaviyo, TikTok Ads, WooCommerce, QuickBooks, banking, 3PLs).',
 'Is cash-flow forecasting live? SKU demand forecasting? What-if scenarios? Counterfactuals on Numbers (and when)?',
 'Which profit metrics are computed today: SKU-level profit, returns and refunds, payment and marketplace fees, contribution margin layers?',
 'Capital: in which countries is it available today, through which named lending partner, and what wording should documentation use?',
 'Trial: how long, is a card required, and what happens at the end?',
 'Renewal and cancellation: which wording is correct?',
 'Which legal entity contracts with customers?',
 'EU data residency: available in practice? For which customers?',
 'Can a customer export all of their data, and in what format?',
 'Can a customer get read access to their own BigQuery dataset?',
 'Support: channels, hours and response time. Is there a named success manager?',
 'Are WhatsApp and email live as ways to talk to Luca?',
 'Does Luca generate board packs / investor reports today?',
 'When Luca is unsure, does it say so or ask a clarifying question? Can a user see how an answer was computed?',
 'Role-based access: can access be limited per data source and per person?',
 'Multi-store and multi-currency: supported? How?',
 'Security documentation: what can be shared with a buyer on request (policies, pen test, DPA, sub-processors)?',
 'For each live connector: what data does Luca pull, and which permission scopes does it request (read-only or write)?',
 'Knowledge base: which file types and sizes are supported, and how long are uploaded files kept?',
 'Which customer brands have approved being named, and which case studies (e.g. the $72,000 subscription-skip finding) can be cited publicly?',
]

thin = Side(style='thin', color='D0D5DD'); B = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR = PatternFill('solid', fgColor='01819A'); FILL = PatternFill('solid', fgColor='FFF7CC'); EX = PatternFill('solid', fgColor='EEF6F8')
F = lambda **k: Font(name='Arial', **{'size': 10, **k})
def sheet(wb, title, headers, data, widths, fill_from, validation=None, example=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for i, h in enumerate(headers, 1):
        c = ws.cell(1, i); c.font = F(bold=True, color='FFFFFF'); c.fill = HDR; c.alignment = Alignment(wrap_text=True, vertical='center'); c.border = B
    rows_ = ([example] if example else []) + data
    for r in rows_: ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            c.font = F(); c.alignment = Alignment(wrap_text=True, vertical='top'); c.border = B
            if c.column >= fill_from: c.fill = FILL
    if example:
        for c in ws[2]: c.fill = EX; c.font = F(italic=True, color='475467')
    for i, w in enumerate(widths, 1): ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(2, 3); ws.auto_filter.ref = ws.dimensions
    if validation:
        dv = DataValidation(type='list', formula1=validation, allow_blank=True); ws.add_data_validation(dv)
        dv.add(f'{get_column_letter(fill_from)}{3 if example else 2}:{get_column_letter(fill_from)}{ws.max_row}')
    return ws

wb = Workbook(); st = wb.active; st.title = 'Start here'
intro = [
 ('Luca AI · product documentation v1 · parameters to confirm', True),
 ('Prepared by MaximusLabs, 16 October 2026', False), ('', False),
 ('Why this exists', True),
 ('Version 1 of Luca\'s product documentation contains only what we are certain Luca has today, from our calls, your product screenshots and your website. Everything we are not certain about is in this workbook. Your answers decide what goes into version 2.', False), ('', False),
 ('How to fill it in', True),
 ('1. Yellow cells are yours. Everything else is our current understanding: correct it freely.', False),
 ('2. Start with tab 1 (website contradictions) and tab 2 (product questions). They unblock the most.', False),
 ('3. On tab 3, work P1 rows first. Pick a status from the dropdown; add plans, limits and a proof link where you can.', False),
 ('4. On tab 4, check what version 1 already says about Luca and correct anything wrong.', False),
 ('5. The first row of each tab (shaded, italic) is an example of the expected format. Leave it as is.', False), ('', False),
 ('Progress', True),
]
for t, b in intro:
    st.append([t]); st.cell(st.max_row, 1).font = F(bold=b, size=14 if st.max_row == 1 else 10)
st.column_dimensions['A'].width = 62; st.column_dimensions['B'].width = 14
for c in st['A']: c.alignment = Alignment(wrap_text=True, vertical='top')

ex1 = ['e.g.', 'Trial length', 'Stated as 14 days and as 30 days.', '/pricing', '30 days, card required, cancel any time before day 30.', 'Eric', 'Pricing page is correct.']
sheet(wb, '1 · Website contradictions', ['#', 'Topic', 'What we see', 'Where', 'Correct version (yours)', 'Confirmed by', 'Notes'],
      [[i+1, *c, '', '', ''] for i, c in enumerate(CONTRA)], [5, 24, 70, 26, 48, 16, 30], 5, example=ex1)
sheet(wb, '2 · Product questions', ['#', 'Question', 'Answer (yours)', 'Confirmed by', 'Proof / link'],
      [[i+1, q, '', '', ''] for i, q in enumerate(QUESTIONS)], [5, 80, 60, 16, 30], 3,
      example=['e.g.', 'Is Klaviyo a live connector?', 'Yes, live since July 2026 on all plans. Read-only.', 'Usman', 'https://… (screenshot)'])
STAT = '"Live,Partial,In build,Planned,Not offered,Not sure"'
sheet(wb, '3 · Parameters to confirm', ['ID', 'Area', 'Parameter', 'Buyer importance', 'Our current understanding', 'Your website', 'Priority', 'What we need', 'Status (yours)', 'Plans / limits', 'Live since', 'Proof (screenshot, doc, link)', 'Correct wording for the docs', 'Notes'],
      [r + ['', '', '', '', '', ''] for r in rows], [7, 22, 40, 11, 22, 24, 8, 34, 14, 26, 12, 28, 40, 24], 9, validation=STAT,
      example=['e.g.', 'Data connectors', 'Klaviyo', 'High', 'Unknown to us', 'Stated on your website', 'P1', 'Please confirm it is live today.', 'Live', 'All plans; read-only', 'Jul 2026', 'https://…', 'Luca connects to Klaviyo (read-only) to analyse flows and campaigns.', ''])
doc = [[p['id'], p['name'] or p['id'], p['dest'], 'Available' if p['status']=='available' else 'Available, limited scope', p['scope'] or '', ' · '.join(f for f, _ in p['facts']), '', ''] for p in spec.P]
sheet(wb, '4 · Documented in v1', ['ID', 'What v1 documents', 'Page', 'Status in v1', 'Scope we state', 'What we based it on', 'Correct? (yours)', 'Correction'],
      doc, [7, 36, 34, 18, 34, 70, 14, 40], 7, validation='"Yes,No,Edit"',
      example=['e.g.', 'Slack app', '/integrations/workspace/slack/', 'Available', '', 'Add Luca to Slack…', 'Edit', 'Also works in Slack Connect channels.'])

tabs = [('1 · Website contradictions', 'E', len(CONTRA)), ('2 · Product questions', 'C', len(QUESTIONS)), ('3 · Parameters to confirm', 'I', len(rows)), ('4 · Documented in v1', 'G', len(doc))]
for name, col, n in tabs:
    r = st.max_row + 1
    st.cell(r, 1, f'{name}: answered').font = F()
    st.cell(r, 2, f"=COUNTA('{name}'!{col}3:{col}{n+2})&\" of {n}\"").font = F(bold=True)
st.append([]); st.append(['Counts exclude the example row.']); st.cell(st.max_row, 1).font = F(italic=True, color='667085')
wb.save(OUT)
print('rows to confirm', len(rows), 'P1', sum(r[6]=='P1' for r in rows), 'P2', sum(r[6]=='P2' for r in rows), 'P3', sum(r[6]=='P3' for r in rows), 'documented', len(doc))
